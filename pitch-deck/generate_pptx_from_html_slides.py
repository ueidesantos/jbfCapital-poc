#!/usr/bin/env python3
"""
JBF Capital - Gerador de PowerPoint dos Slides HTML
Gera apresentação PowerPoint usando os screenshots capturados dos slides HTML
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
import os
from pathlib import Path

# Cores da Identidade Visual JBF Capital
class JBFColors:
    PRIMARY = RGBColor(102, 126, 234)  # #667EEA
    SECONDARY = RGBColor(118, 75, 162)  # #764BA2
    DARK_SLATE = RGBColor(30, 41, 59)   # #1E293B
    SLATE_GRAY = RGBColor(100, 116, 139)  # #64748B
    WHITE = RGBColor(255, 255, 255)

def create_presentation_from_html_slides():
    """
    Cria apresentação PowerPoint usando os screenshots dos slides HTML
    """
    print("=" * 70)
    print("JBF CAPITAL - GERADOR DE POWERPOINT DOS SLIDES HTML")
    print("=" * 70)
    print()
    
    # Criar apresentação com proporção 16:9
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(5.625)
    
    # Diretório com os screenshots
    images_dir = Path("slides_from_html")
    
    # Verificar se o diretório existe
    if not images_dir.exists():
        print(f"❌ Erro: Diretório '{images_dir}' não encontrado!")
        print("Execute primeiro: python3 capture_html_slides.py")
        return False
    
    # Listar todas as imagens PNG no diretório
    image_files = sorted([f for f in images_dir.glob("Slide_*.png")])
    
    if not image_files:
        print(f"❌ Erro: Nenhuma imagem Slide_*.png encontrada em '{images_dir}'!")
        print("Execute primeiro: python3 capture_html_slides.py")
        return False
    
    print(f"Encontrados {len(image_files)} slides:")
    for img in image_files:
        print(f"  • {img.name}")
    print()
    
    # Criar um slide para cada imagem
    for idx, image_file in enumerate(image_files, 1):
        print(f"Adicionando Slide {idx}: {image_file.name}...")
        
        # Adicionar slide em branco
        blank_slide_layout = prs.slide_layouts[6]  # Layout em branco
        slide = prs.slides.add_slide(blank_slide_layout)
        
        # Dimensões para slide completo (sem bordas)
        # Usar exatamente o tamanho do slide
        img_width = prs.slide_width
        img_height = prs.slide_height
        left = Inches(0)
        top = Inches(0)
        
        # Adicionar imagem ao slide (tela cheia)
        try:
            slide.shapes.add_picture(
                str(image_file),
                left,
                top,
                width=img_width,
                height=img_height
            )
        except Exception as e:
            print(f"  ⚠️  Erro ao adicionar imagem {image_file.name}: {e}")
            continue
    
    # Salvar apresentação
    output_file = "JBF_Capital_Pitch_Deck.pptx"
    prs.save(output_file)
    
    print()
    print("=" * 70)
    print(f"✅ Apresentação gerada com sucesso!")
    print(f"📊 Arquivo: {Path(output_file).resolve()}")
    print(f"📸 Total de slides: {len(image_files)}")
    print(f"📦 Tamanho: {os.path.getsize(output_file) / 1024:.0f} KB")
    print("=" * 70)
    print()
    print("A apresentação está pronta para uso!")
    print("Abra o arquivo no PowerPoint, Keynote ou Google Slides.")
    print()
    
    return True

if __name__ == "__main__":
    # Mudar para o diretório pitch-deck
    script_dir = Path(__file__).parent
    os.chdir(script_dir)
    
    success = create_presentation_from_html_slides()
    
    if not success:
        import sys
        sys.exit(1)
